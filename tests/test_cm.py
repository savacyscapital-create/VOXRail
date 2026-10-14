"""
Tests for the voxrail_cm configuration management plugin.
Run from the project root:  GROQ_API_KEY=dummy python -m pytest tests -q
The LLM is stubbed; responses below follow the VoxRail templates.
"""

import os
import sys

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ.setdefault("GROQ_API_KEY", "dummy")

UC1 = """**Use Case No**: UC-MP-001
**Use Case Title**: Open passenger doors at platform
**Use Case Description**: Driver releases doors on the platform side.
**Previous Train / System State**: Stopped at Platform
**New Train / System State**: Doors Released
**Applicable Train States**: Stopped at Platform, Ready to Drive
**State Transition Expected Time**: 2 s
**Primary Actors**: Train Driver, Passenger
**Logical Interfaces**: Door Release Command, Door Closed Loop
**Trigger**: Driver presses door release
"""

UC2 = """**Use Case No**: UC-MP-002
**Use Case Title**: Traction enable
**Previous Train / System State**: Ready-to-Drive
**New Train / System State**: Traction Enabled
**Applicable Train States**: Ready to Drive
**Primary Actors**: Driver
**Logical Interfaces**: Door Closed Loop
"""

UC3 = """**Use Case No**: UC-MP-003
**Use Case Title**: Cab heating
**Applicable Train States**: Stabled
**Primary Actors**: Maintainer
"""

TC1 = """**Test Case Ref**: TC-MP-001
**Use Case Ref**: UC-MP-001
**Use Case Title**: Open passenger doors at platform
**Test Description**: Verify door release
"""

ED1 = """⚠️ DESIGN NOTICE: logical interfaces only.

**Electrical Diagram Ref**: ED-MP-001
**Related UC**: UC-MP-001
**System / Function**: Door control
**Train States Served**: Stopped at Platform, Ready to Drive
**Logical Interfaces**: Door Closed Loop, Discrete, 110V DC; Door Release Command, MVB
**Voltage Domain**: 110V DC

```mermaid
flowchart LR
  DCU[Door Control Unit] -->|Door Closed Loop, Discrete, 110V DC| TCMS[TCMS]
  TCMS -->|Door Release Command, MVB| DCU
```
"""

SD1 = """**State Diagram Ref**: SD-MP-001
**Related UC**: UC-MP-001
**Train States Modelled**: Stopped at Platform

```mermaid
stateDiagram-v2
  [*] --> Stopped_at_Platform
  Stopped_at_Platform --> Doors_Released : DoorReleaseCmd
  Doors_Released --> Ready_to_Drive : DoorsClosedLocked
```
"""

FM1 = """⚠️ SAFETY NOTICE: AI-generated.

**FMEA Ref**: FM-MP-001
**Related UC**: UC-MP-001
**Related State**: Doors Released
**Failure Mode**: Door Closed Loop falsely reports closed
"""

HZ1 = """**Hazard Ref**: HZ-MP-001
**Related UC**: UC-MP-001
**Related FMEA Ref**: FM-MP-001
**Hazard Description**: Train departs with door open
"""

CC1 = """**Code Ref**: CC-MP-001
**Related UC**: UC-MP-001
**Related State Diagram**: SD-MP-001
"""


@pytest.fixture()
def env(tmp_path, monkeypatch):
    monkeypatch.setenv("VOXRAIL_CM_DIR", str(tmp_path))
    from voxrail_cm import store as cm_store
    cm_store.reset_cache()
    import main
    yield main
    cm_store.reset_cache()


def feed(main, session, *responses):
    for r in responses:
        session.update_from_response(r)
        main.emit("on_response", session, r)


def test_plugin_loaded(env):
    assert any(getattr(p, "name", "") == "config_management" for p in env.PLUGINS)


def test_parser_finds_fields():
    from voxrail_cm.parser import parse_response
    [uc] = parse_response(UC1)
    assert uc.ref == "UC-MP-001"
    assert uc.actors == ["Train Driver", "Passenger"]
    assert uc.states == ["Stopped at Platform", "Ready to Drive"]
    assert uc.from_states == ["Stopped at Platform"] and uc.to_states == ["Doors Released"]
    [ed] = parse_response(ED1)
    assert [i["name"] for i in ed.interfaces] == ["Door Closed Loop", "Door Release Command"]
    assert ed.interfaces[1]["kind"] == "MVB"
    [sd] = parse_response(SD1)
    names = {e["name"] for e in sd.sw_elements}
    assert {"Stopped at Platform", "Doors Released", "DoorReleaseCmd"} <= names


def test_two_artefacts_in_one_response():
    from voxrail_cm.parser import parse_response
    arts = parse_response(UC3 + "\n\n" + TC1)
    assert [a.ref for a in arts] == ["UC-MP-003", "TC-MP-001"]


def test_placeholders_ignored():
    from voxrail_cm.parser import split_list
    assert split_list("[comma-separated; every train state]") == []
    assert split_list("N/A") == []


def test_layer1_coupling_through_state_and_interface(env):
    s = env.VoxRailSession("Metro Project")
    feed(env, s, UC1, UC2, UC3)
    from voxrail_cm import graph as g
    from voxrail_cm.store import get_store
    st = get_store(s.project_code)
    c = g.coupled_use_cases(st, "UC-MP-001")
    coupled = {x["uc"]["key"]: {v["via"]["label"] for v in x["through"]} for x in c["coupled"]}
    assert set(coupled) == {"UC-MP-002"}                       # UC-003 shares nothing
    assert {"Ready to Drive", "Door Closed Loop"} <= coupled["UC-MP-002"]


def test_near_duplicate_names_suggested_and_merged(env):
    s = env.VoxRailSession("Metro Project")
    feed(env, s, UC1, UC2)
    from voxrail_cm.store import get_store
    st = get_store(s.project_code)
    # "Ready-to-Drive" normalises to "ready to drive": same entity, no suggestion
    assert len([e for e in st.entities("state") if e["name"].lower().startswith("ready")]) == 1
    # "Driver" vs "Train Driver" is a near match: suggested, not merged silently
    sugg = st.suggestions()
    assert any({x["new_name"], x["candidate_name"]} == {"Driver", "Train Driver"} for x in sugg)
    sid = next(x["id"] for x in sugg if x["kind"] == "actor")
    st.resolve_suggestion(sid, accept=True)
    actors = [e["name"] for e in st.entities("actor")]
    assert actors.count("Driver") + actors.count("Train Driver") == 1
    # UC-002 now shares the actor with UC-001
    from voxrail_cm import graph as g
    hubs = {h["kind"] for x in g.coupled_use_cases(st, "UC-MP-002")["coupled"] for h in
            [v["via"] for v in x["through"]]}
    assert "actor" in hubs


def test_layer2_state_view(env):
    s = env.VoxRailSession("Metro Project")
    feed(env, s, UC1, UC2, ED1, SD1, TC1, FM1, HZ1, CC1)
    from voxrail_cm import graph as g
    from voxrail_cm.store import get_store, ekey
    st = get_store(s.project_code)
    rtd = st.find_entity("state", "Ready to Drive")
    v = g.state_view(st, ekey("state", rtd["id"]))
    ucs = {n["key"] for n in v["use_case_level"]["valid_in"]}
    assert ucs == {"UC-MP-001", "UC-MP-002"}
    ifaces = {n["label"] for n in v["logical_level"]["realised_by_electrical"]}
    assert ifaces == {"Door Closed Loop", "Door Release Command"}
    # Stopped at Platform shares the ED interfaces -> coupled at logical level
    assert any(c["state"]["label"] == "Stopped at Platform" for c in v["coupled_states"])
    sap = st.find_entity("state", "Stopped at Platform")
    v2 = g.state_view(st, ekey("state", sap["id"]))
    assert v2["logical_level"]["realised_by_software"]
    # traceability links from TC/FM/HZ/CC
    a = st.artefact("UC-MP-001")
    kinds = {l["kind"] for l in a["links_in"]}
    assert {"verifies", "analyses", "derived_from", "realises", "models_uc", "implements"} <= kinds


def test_regeneration_versions_and_relinks(env):
    s = env.VoxRailSession("Metro Project")
    feed(env, s, UC1)
    changed = UC1.replace("Stopped at Platform, Ready to Drive", "Stopped at Platform")
    feed(env, s, changed)
    from voxrail_cm.store import get_store
    st = get_store(s.project_code)
    a = st.artefact("UC-MP-001")
    assert a["version"] == 2
    states = {st.label(l["dst"])["label"] for l in a["links_out"] if l["kind"] == "valid_in"}
    assert states == {"Stopped at Platform"}


def test_gaps(env):
    s = env.VoxRailSession("Metro Project")
    feed(env, s, UC1, UC2, TC1, HZ1)
    from voxrail_cm import graph as g
    from voxrail_cm.store import get_store
    gp = g.gaps(get_store(s.project_code))
    assert gp["uc_without_test"] == ["UC-MP-002"]
    assert {"from": "HZ-MP-001", "to": "FM-MP-001", "kind": "derived_from"} in gp["dangling_refs"]


def test_prompt_context_and_no_db_for_plain_chat(env, tmp_path):
    s = env.VoxRailSession("Probe")
    feed(env, s, "OK")
    assert not list(tmp_path.glob("*.db"))                      # plain replies create no graph
    assert "=== CONFIGURATION CATALOGUE" not in env.build_system_prompt(s)
    feed(env, s, UC1.replace("UC-MP-", "UC-P-"))
    p = env.build_system_prompt(s)
    assert "Train States: " in p and "Ready to Drive" in p
    assert "**Applicable Train States**" in p and "**State Diagram Ref**: SD-P-001" in p


def test_chat_path_with_stub_llm(env, monkeypatch):
    class Msg:
        def __init__(self, c): self.content = c
    monkeypatch.setattr(env.llm.__class__, "invoke", lambda self, msgs: Msg(SD1.replace("-MP-", "-MP-")))
    s = env.VoxRailSession("Metro Project")
    env.chat(s, "state diagram please")
    assert s.sd_counter == 1 and s.peek_sd_ref() == "SD-MP-002"
    from voxrail_cm.store import get_store
    assert get_store("MP").artefact("SD-MP-001")


def test_api(env):
    from fastapi.testclient import TestClient
    import app as webapp
    s = env.VoxRailSession("Metro Project")
    feed(env, s, UC1, UC2, ED1, SD1, TC1)
    c = TestClient(webapp.app)
    q = {"project_code": "MP"}
    assert c.get("/cm/summary", params=q).json()["artefacts"]["UC"] == 2
    assert c.get("/cm/graph", params={**q, "layer": "1"}).status_code == 200
    l2 = c.get("/cm/graph", params={**q, "layer": "2"}).json()
    assert all(n["layer"] in (0, 2) for n in l2["nodes"])
    assert "flowchart LR" in c.get("/cm/graph.mmd", params=q).text
    html = c.get("/cm/view", params=q).text
    assert "mermaid.render" in html and "</script>" not in html.split("const SRC")[1].split(";\n")[0]
    st = c.get("/cm/states/Ready to Drive", params=q).json()
    assert st["use_case_count"] == 2
    assert c.get("/cm/uc/UC-MP-001/coupled", params=q).json()["coupled"]
    assert c.post("/cm/states", params=q, json={"names": ["Stabled", "Coupled"]}).status_code == 200
    assert c.get("/cm/states/Nope", params=q).status_code == 404
    assert c.get("/cm/summary").status_code == 404                 # no session, no project_code
    assert "config_management" in c.get("/health").json()["plugins"]


# ── Matrix views ──────────────────────────────────────────────────────────────
def _full_store(env):
    s = env.VoxRailSession("Metro Project")
    feed(env, s, UC1, UC2, UC3, ED1, SD1, TC1, FM1, HZ1, CC1)
    from voxrail_cm.store import get_store
    return get_store(s.project_code)


def _row(m, key):
    return next(r for r in m["rows"] if r["key"] == key or r["label"] == key)


def _col(m, label):
    return next(c["key"] for c in m["columns"] if c["label"] == label)


def test_matrix_uc_state(env):
    from voxrail_cm import matrix as mx
    m = mx.build(_full_store(env), "uc_state")
    r1 = _row(m, "UC-MP-001")
    assert r1["cells"][_col(m, "Stopped at Platform")] == "VF"
    assert r1["cells"][_col(m, "Doors Released")] == "T"
    assert r1["fields"]["Actors"] == "Passenger, Train Driver"
    # filters
    assert [r["key"] for r in mx.build(_full_store(env), "uc_state", state="stabled")["rows"]] == ["UC-MP-003"]
    assert [r["key"] for r in mx.build(_full_store(env), "uc_state", actor="maint")["rows"]] == ["UC-MP-003"]
    m2 = mx.build(_full_store(env), "uc_state", col="ready", hide_empty=True)
    assert [c["label"] for c in m2["columns"]] == ["Ready to Drive"]
    assert {r["key"] for r in m2["rows"]} == {"UC-MP-001", "UC-MP-002"}


def test_matrix_state_logic(env):
    from voxrail_cm import matrix as mx
    m = mx.build(_full_store(env), "state_logic")
    sap = _row(m, "Stopped at Platform")
    assert sap["cells"][_col(m, "Door Closed Loop")] == "E"
    assert "S" in sap["cells"].values()
    assert not sap["gap"]
    stabled = _row(m, "Stabled")
    assert stabled["gap"] and "no electrical logic" in stabled["gap_note"]
    assert {c["group"] for c in m["columns"]} == {"Electrical interface", "Software state"}


def test_matrix_trace_and_gaps(env):
    from voxrail_cm import matrix as mx
    m = mx.build(_full_store(env), "uc_trace")
    r1 = _row(m, "UC-MP-001")
    assert r1["cells"] == {"TC": "TC-MP-001", "FM": "FM-MP-001", "HZ": "HZ-MP-001",
                           "ED": "ED-MP-001", "SD": "SD-MP-001", "CC": "CC-MP-001"}
    assert not r1["gap"]
    g = mx.build(_full_store(env), "uc_trace", gaps_only=True)
    assert {r["key"] for r in g["rows"]} == {"UC-MP-002", "UC-MP-003"}
    assert all("no test" in r["gap_note"] for r in g["rows"])


def test_matrix_uc_uc(env):
    from voxrail_cm import matrix as mx
    m = mx.build(_full_store(env), "uc_uc")
    cell = _row(m, "UC-MP-001")["cells"]["UC-MP-002"]
    assert "S:Ready to Drive" in cell and "I:Door Closed Loop" in cell
    assert "UC-MP-003" not in _row(m, "UC-MP-001")["cells"]
    assert _row(m, "UC-MP-003")["fields"]["Coupled with"] == 0


def test_matrix_exports(env):
    import io as _io
    from openpyxl import load_workbook
    from voxrail_cm import matrix as mx
    st = _full_store(env)
    csv_text = mx.to_csv(mx.build(st, "uc_trace"))
    assert csv_text.splitlines()[0].startswith("Key,Title,States,Version,Tests,FMEA")
    wb = load_workbook(_io.BytesIO(mx.to_xlsx([mx.build(st, n) for n in mx.MATRICES])))
    assert wb.sheetnames == list(mx.MATRICES)
    assert wb["uc_state"].auto_filter.ref.startswith("A1:")


def test_matrix_api(env):
    from fastapi.testclient import TestClient
    import app as webapp
    _full_store(env)
    c = TestClient(webapp.app)
    q = {"project_code": "MP"}
    assert len(c.get("/cm/matrix").json()) == 4
    j = c.get("/cm/matrix/uc_trace", params={**q, "gaps_only": "true"}).json()
    assert j["totals"]["gaps"] == 2
    r = c.get("/cm/matrix/uc_state", params={**q, "format": "csv"})
    assert r.headers["content-type"].startswith("text/csv") and "UC-MP-001" in r.text
    x = c.get("/cm/matrix.xlsx", params=q)
    assert x.status_code == 200 and x.content[:2] == b"PK"
    assert c.get("/cm/matrix/nope", params=q).status_code == 404
    assert "VoxRail configuration matrices" in c.get("/cm/matrices").text


# ── Phase 2: change management ────────────────────────────────────────────────
FM_CRIT = FM1.replace("**Failure Mode**:", "**Severity**: Critical\n**Failure Mode**:")


def _cm_store(env, *extra):
    s = env.VoxRailSession("Metro Project")
    feed(env, s, UC1, UC2, UC3, ED1, SD1, TC1, FM_CRIT, HZ1, CC1, *extra)
    from voxrail_cm.store import get_store
    return get_store(s.project_code)


def _cats(imp):
    return {i["key"] if i["kind"] in ("UC", "TC", "FM", "HZ", "ED", "SD", "CC") else f'{i["kind"]}:{i["label"]}':
            i["category"] for i in imp["items"]}


def test_impact_of_use_case_change(env):
    from voxrail_cm import change as ch
    st = _cm_store(env)
    imp = ch.impact(st, ["UC-MP-002"])
    c = _cats(imp)
    assert c["UC-MP-002"] == "Changed"
    assert c["UC-MP-001"] == "Coupled via state"           # shares Ready to Drive
    assert "UC-MP-003" not in c                             # Stabled only: out of scope
    assert c["TC-MP-001"] == "Indirect dependent"           # test of the coupled UC
    assert c["interface:Door Closed Loop"] == "Logical design"
    assert c["state:Ready to Drive"] == "Coupled via state"
    # actor coupling is opt-in; UC-003's actor (Maintainer) is not shared, so it stays out
    assert "UC-MP-003" not in _cats(ch.impact(st, ["UC-MP-002"], include_actor_coupling=True))


def test_impact_of_interface_change_reaches_use_cases_through_states(env):
    from voxrail_cm import change as ch
    st = _cm_store(env)
    key = ch.resolve_target(st, "interface:Door Release Command")
    c = _cats(ch.impact(st, [key]))
    assert c["interface:Door Release Command"] == "Changed"
    assert c["ED-MP-001"] == "Direct dependent"
    assert c["state:Stopped at Platform"] == "Direct dependent"
    assert c["UC-MP-001"] == "Coupled via state"
    assert c["UC-MP-002"] == "Coupled via state"            # Ready to Drive is served by the same interface
    assert "UC-MP-003" not in c


def test_impact_of_state_change_and_safety_flags(env):
    from voxrail_cm import change as ch
    st = _cm_store(env)
    imp = ch.impact(st, [ch.resolve_target(st, "state:Stopped at Platform")])
    items = {i["key"]: i for i in imp["items"]}
    assert items["FM-MP-001"]["severity"] == "Critical" and items["FM-MP-001"]["safety"]
    assert items["HZ-MP-001"]["safety"]
    assert items["UC-MP-001"]["safety"]
    assert set(imp["safety_items"]) >= {"FM-MP-001", "HZ-MP-001", "UC-MP-001"}


def test_verification_scope(env):
    from voxrail_cm import change as ch
    st = _cm_store(env)
    imp = ch.impact(st, ["UC-MP-002"])
    sc = ch.verification_scope(st, imp)
    rows = {r["key"]: r for r in sc["rows"]}
    assert rows["UC-MP-001"]["cells"]["rerun_tests"] == "TC-MP-001"
    assert rows["UC-MP-001"]["cells"]["hazard"] == "HZ-MP-001"
    assert rows["UC-MP-002"]["cells"]["new_test"].startswith("Yes")
    assert sc["summary"]["tests_to_rerun"] == ["TC-MP-001"]
    assert sc["summary"]["new_tests_needed"] == ["UC-MP-002"]
    assert [o["key"] for o in sc["out_of_scope"]] == ["UC-MP-003"]


def test_cr_lifecycle_and_frozen_assessment(env):
    from voxrail_cm import change as ch
    st = _cm_store(env)
    cr = ch.create_cr(st, "MP", "Change door release logic", targets=[{"target": "ED-MP-001"}])
    assert cr["code"] == "CR-MP-001" and cr["status"] == "Open"
    with pytest.raises(ValueError):
        ch.add_target(st, "CR-MP-001", "state:Nowhere")
    ch.add_target(st, "CR-MP-001", "state:Ready to Drive")
    res = ch.assess(st, "CR-MP-001")
    assert res["cr"]["status"] == "Assessed"
    a = ch.last_assessment(st, "CR-MP-001")
    assert a["artefacts_changed_since"] == []
    # regenerate UC-001 after assessment -> assessment reports it changed
    s = env.VoxRailSession("Metro Project")
    feed(env, s, UC1.replace("2 s", "3 s"))
    assert ch.last_assessment(st, "CR-MP-001")["artefacts_changed_since"] == ["UC-MP-001"]
    ch.set_status(st, "CR-MP-001", "Approved", "CCB 14 Oct")
    assert any("Approved" in h["event"] for h in ch.get_cr(st, "CR-MP-001")["history"])


def test_variation_order_creates_draft_cr(env):
    from voxrail_cm import change as ch
    vo = "⚠️ VARIATION ORDER DETECTED: Add cab heating pre-conditioning\n\n" + UC3
    st = _cm_store(env, vo)
    crs = ch.list_crs(st)
    assert len(crs) == 1 and crs[0]["status"] == "Draft" and crs[0]["origin"] == "variation_order"
    assert crs[0]["title"] == "Add cab heating pre-conditioning"
    assert [t["key"] for t in crs[0]["targets"]] == ["UC-MP-003"]


def test_change_api(env):
    import io as _io
    from fastapi.testclient import TestClient
    from openpyxl import load_workbook
    import app as webapp
    _cm_store(env)
    c = TestClient(webapp.app)
    q = {"project_code": "MP"}
    assert c.post("/cm/cr", params=q, json={"title": "x", "targets": [{"target": "UC-XX-999"}]}).status_code == 400
    cr = c.post("/cm/cr", params=q, json={"title": "Traction enable timing",
                                          "targets": [{"target": "UC-MP-002"}]}).json()
    code = cr["code"]
    imp = c.get(f"/cm/cr/{code}/impact", params=q).json()
    assert any(i["key"] == "UC-MP-001" for i in imp["items"])
    assert "Coupled via state" in c.get(f"/cm/cr/{code}/impact", params={**q, "format": "csv"}).text
    scope = c.get(f"/cm/cr/{code}/scope", params=q).json()
    assert scope["summary"]["tests_to_rerun"] == ["TC-MP-001"]
    assert c.get(f"/cm/cr/{code}/assessment", params=q).status_code == 404
    assert c.post(f"/cm/cr/{code}/assess", params=q).json()["cr"]["status"] == "Assessed"
    x = c.get(f"/cm/cr/{code}/report.xlsx", params={**q, "frozen": "true"})
    assert load_workbook(_io.BytesIO(x.content)).sheetnames == ["impact", "verification_scope", "out_of_scope"]
    assert c.post(f"/cm/cr/{code}/status", params=q, json={"status": "Bogus"}).status_code == 400
    assert c.get("/cm/cr/CR-MP-999", params=q).status_code == 404
    wi = c.get("/cm/impact", params=[("project_code", "MP"), ("target", "state:Stabled")]).json()
    assert [i["key"] for i in wi["impact"]["items"] if i["kind"] == "UC"] == ["UC-MP-003"]
    assert len(c.get("/cm/cr", params=q).json()) == 1
    assert "VoxRail change management" in c.get("/cm/changes").text


# ── Unique project codes ──────────────────────────────────────────────────────
def test_same_initials_get_unique_codes(env):
    a = env.VoxRailSession("Metro Project")
    b = env.VoxRailSession("Mainline Programme")
    c = env.VoxRailSession("Main Platform")
    assert (a.project_code, b.project_code, c.project_code) == ("MP", "MP2", "MP3")
    # same name (any case / spacing) resumes the same project
    assert env.VoxRailSession("  metro   PROJECT ").project_code == "MP"


def test_counters_survive_restart_and_tabs(env):
    s = env.VoxRailSession("Metro Project")
    s.update_from_response(UC1)
    s.update_from_response(UC2)
    assert s.peek_uc_ref() == "UC-MP-003"
    again = env.VoxRailSession("Metro Project")          # server restart / new tab
    assert again.peek_uc_ref() == "UC-MP-003"
    again.update_from_response(UC3)                       # tab 2 creates UC-MP-003
    s.sync_counters()                                     # tab 1 catches up before its next chat
    assert s.peek_uc_ref() == "UC-MP-004"


def test_probe_session_not_registered(env, tmp_path):
    env.VoxRailSession("probe", register=False)
    assert not (tmp_path / "projects.json").exists()


def test_projects_endpoint(env):
    from fastapi.testclient import TestClient
    import app as webapp
    env.VoxRailSession("Metro Project")
    env.VoxRailSession("Mainline Programme")
    codes = [p["project_code"] for p in TestClient(webapp.app).get("/projects").json()]
    assert codes == ["MP", "MP2"]
