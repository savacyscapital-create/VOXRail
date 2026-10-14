# =============================================================================
# VoxRail — Railway Systems Engineering AI
# Copyright © 2026 Savacys Capital. All rights reserved.
#
# Licensed under the GNU Affero General Public License v3.0 (AGPL-3.0)
# for non-commercial use. See LICENSE for full terms.
# =============================================================================
"""
voxrail_cm.matrix
=================
Matrix views over the configuration graph, for filtering, analysis and export.

  uc_state     Use case × Train state           (V = valid in, F = from, T = to)
  state_logic  Train state × Interface / SW     (E = electrical, S = software)
  uc_trace     Use case × TC / FM / HZ / ED / SD / CC   (traceability + coverage gaps)
  uc_uc        Use case × Use case              (shared states / actors / interfaces)

Every matrix has the same shape so the page, CSV and XLSX export share one code path:
  {name, title, legend, row_fields: [..], columns: [{key, label, group}],
   rows: [{key, label, fields: {..}, cells: {col_key: text}, gap: bool, gap_cols: [..]}]}
"""

import csv
import io
from collections import defaultdict

from .store import ConfigStore, UC_STATE_LINKS

MATRICES = {
    "uc_state":    "Use case × Train state",
    "state_logic": "Train state × Interface / software logic",
    "uc_trace":    "Use case traceability (TC / FMEA / Hazard / design)",
    "uc_uc":       "Use case × Use case coupling",
}
STATE_CODE = {"valid_in": "V", "transitions_from": "F", "transitions_to": "T"}
TRACE_COLS = [("TC", "verifies", "Tests"), ("FM", "analyses", "FMEA"), ("HZ", "derived_from", "Hazards"),
              ("ED", "realises", "Electrical"), ("SD", "models_uc", "State diagrams"),
              ("CC", "implements", "C code")]


# ── Shared lookups ────────────────────────────────────────────────────────────
class _Ctx:
    def __init__(self, store: ConfigStore):
        self.store = store
        self.links = store.links()
        self.ents = {e["key"]: e for e in store.entities()}
        self.arts = {a["ref"]: a for a in store.artefacts()}
        self.out = defaultdict(list)
        self.inc = defaultdict(list)
        for l in self.links:
            self.out[l["src"]].append(l)
            self.inc[l["dst"]].append(l)

    def name(self, key: str) -> str:
        e = self.ents.get(key)
        return e["name"] if e else key

    def ucs(self) -> list[str]:
        refs = {r for r, a in self.arts.items() if a["type"] == "UC"}
        refs |= {l["dst"] for l in self.links if l["dst"].startswith("UC-")}   # referenced but not generated
        return sorted(refs)

    def uc_hubs(self, uc: str, prefix: str) -> list[str]:
        return [l["dst"] for l in self.out[uc] if l["dst"].startswith(prefix + ":")]

    def uc_fields(self, uc: str) -> dict:
        a = self.arts.get(uc)
        return {
            "Title":  a["title"] if a else "(not generated)",
            "Actors": ", ".join(sorted({self.name(k) for k in self.uc_hubs(uc, "actor")})),
            "States": ", ".join(sorted({self.name(k) for k in self.uc_hubs(uc, "state")})),
            "Version": a["version"] if a else "",
        }


# ── Builders ──────────────────────────────────────────────────────────────────
def _uc_state(c: _Ctx) -> dict:
    states = sorted((k for k, e in c.ents.items() if e["kind"] == "state"), key=c.name)
    rows = []
    for uc in c.ucs():
        cells = {}
        for l in c.out[uc]:
            if l["kind"] in STATE_CODE:
                cells.setdefault(l["dst"], set()).add(STATE_CODE[l["kind"]])
        cells = {k: "".join(x for x in "VFT" if x in v) for k, v in cells.items()}
        rows.append({"key": uc, "label": uc, "fields": c.uc_fields(uc), "cells": cells,
                     "gap": not cells, "gap_note": "no train state" if not cells else ""})
    return {"row_fields": ["Title", "Actors", "Version"],
            "columns": [{"key": s, "label": c.name(s), "group": "Train state"} for s in states],
            "legend": "V = valid in · F = transitions from · T = transitions to",
            "rows": rows}


def _state_logic(c: _Ctx) -> dict:
    states = sorted((k for k, e in c.ents.items() if e["kind"] == "state"), key=c.name)
    cols = sorted((k for k, e in c.ents.items() if e["kind"] == "interface"), key=c.name)
    sw = sorted((k for k, e in c.ents.items()
                 if e["kind"] == "sw_element" and e["attrs"].get("kind") == "state"),
                key=lambda k: (c.ents[k]["attrs"].get("sd", ""), c.name(k)))
    columns = [{"key": k, "label": c.name(k), "group": "Electrical interface",
                "detail": c.ents[k]["attrs"].get("kind", "")} for k in cols]
    columns += [{"key": k, "label": c.name(k), "group": "Software state",
                 "detail": c.ents[k]["attrs"].get("sd", "")} for k in sw]
    rows = []
    for s in states:
        cells = {}
        for l in c.out[s]:
            if l["kind"] == "realised_by_electrical":
                cells[l["dst"]] = "E"
            elif l["kind"] == "realised_by_software":
                cells[l["dst"]] = "S"
        designed = sorted({l["src"] for l in c.inc[s] if l["kind"] in ("serves_state", "models_state")})
        n_uc = len({l["src"] for l in c.inc[s] if l["kind"] in UC_STATE_LINKS})
        has_e = "E" in cells.values()
        has_s = "S" in cells.values()
        note = ", ".join(x for x, ok in (("no electrical logic", has_e), ("no software logic", has_s)) if not ok)
        rows.append({"key": s, "label": c.name(s),
                     "fields": {"Status": c.ents[s]["status"], "Use cases": n_uc,
                                "Designed in": ", ".join(designed)},
                     "cells": cells, "gap": bool(note), "gap_note": note})
    return {"row_fields": ["Status", "Use cases", "Designed in"], "columns": columns,
            "legend": "E = realised by electrical interface (ED) · S = realised by software state (SD)",
            "rows": rows}


def _uc_trace(c: _Ctx) -> dict:
    rows = []
    for uc in c.ucs():
        cells, gap_cols = {}, []
        for t, kind, _ in TRACE_COLS:
            refs = sorted({l["src"] for l in c.inc[uc] if l["kind"] == kind and l["src"].startswith(t + "-")})
            if refs:
                cells[t] = ", ".join(refs)
        if "TC" not in cells:
            gap_cols.append("TC")
        notes = []
        if "TC" not in cells:
            notes.append("no test")
        if uc not in c.arts:
            notes.append("referenced but not generated")
        if not c.uc_hubs(uc, "state"):
            notes.append("no train state")
        rows.append({"key": uc, "label": uc, "fields": c.uc_fields(uc), "cells": cells,
                     "gap": bool(notes), "gap_cols": gap_cols, "gap_note": ", ".join(notes)})
    return {"row_fields": ["Title", "States", "Version"],
            "columns": [{"key": t, "label": lbl, "group": "Linked artefacts"} for t, _, lbl in TRACE_COLS],
            "legend": "Cells list the artefacts linked to each use case. Gaps: no test, no train state, "
                      "or referenced but never generated.",
            "rows": rows}


def _uc_uc(c: _Ctx) -> dict:
    ucs = c.ucs()
    hubs = {uc: set(c.uc_hubs(uc, "state") + c.uc_hubs(uc, "actor") + c.uc_hubs(uc, "interface"))
            for uc in ucs}
    tag = {"state": "S", "actor": "A", "interface": "I"}
    rows = []
    for a in ucs:
        cells = {}
        for b in ucs:
            if a == b:
                continue
            shared = sorted(hubs[a] & hubs[b], key=lambda k: (k.split(":")[0] != "state", c.name(k)))
            if shared:
                cells[b] = "; ".join(f"{tag[k.split(':')[0]]}:{c.name(k)}" for k in shared)
        f = c.uc_fields(a)
        rows.append({"key": a, "label": a, "fields": {"Title": f["Title"], "Coupled with": len(cells)},
                     "cells": cells, "gap": False, "gap_note": ""})
    return {"row_fields": ["Title", "Coupled with"],
            "columns": [{"key": u, "label": u, "group": "Use case"} for u in ucs],
            "legend": "Shared items: S = train state · A = actor · I = logical interface",
            "rows": rows}


BUILDERS = {"uc_state": _uc_state, "state_logic": _state_logic, "uc_trace": _uc_trace, "uc_uc": _uc_uc}


# ── Public API ────────────────────────────────────────────────────────────────
def build(store: ConfigStore, name: str, row: str | None = None, col: str | None = None,
          actor: str | None = None, state: str | None = None, gaps_only: bool = False,
          hide_empty: bool = False) -> dict:
    """Build one matrix and apply filters.

    row / col   case-insensitive substring on the row (label + fields) / column label
    actor       keep UC rows performed by an actor whose name contains this
    state       keep rows linked to a train state whose name contains this
                (UC rows: via the UC's states; state rows: by name)
    gaps_only   keep rows flagged as gaps
    hide_empty  drop rows and columns with no cells after filtering
    """
    if name not in BUILDERS:
        raise KeyError(name)
    c = _Ctx(store)
    m = BUILDERS[name](c)
    rows, cols = m["rows"], m["columns"]

    def has(text, needle):
        return needle.lower() in str(text).lower()

    if row:
        rows = [r for r in rows if has(r["label"], row) or any(has(v, row) for v in r["fields"].values())]
    if actor and name != "state_logic":
        rows = [r for r in rows if any(has(c.name(k), actor) for k in c.uc_hubs(r["key"], "actor"))]
    if state:
        if name == "state_logic":
            rows = [r for r in rows if has(r["label"], state)]
        else:
            rows = [r for r in rows if any(has(c.name(k), state) for k in c.uc_hubs(r["key"], "state"))]
    if gaps_only:
        rows = [r for r in rows if r["gap"]]
    if col:
        cols = [x for x in cols if has(x["label"], col) or has(x["group"], col)]
    if hide_empty:
        used = {k for r in rows for k in r["cells"]}
        cols = [x for x in cols if x["key"] in used]
        keep = {x["key"] for x in cols}
        rows = [r for r in rows if any(k in keep for k in r["cells"])]
    keep = {x["key"] for x in cols}
    rows = [dict(r, cells={k: v for k, v in r["cells"].items() if k in keep}) for r in rows]

    return {"name": name, "title": MATRICES[name], **m, "columns": cols, "rows": rows,
            "totals": {"rows": len(rows), "columns": len(cols), "gaps": sum(1 for r in rows if r["gap"])}}


def _table(m: dict) -> list[list]:
    head = ["Key"] + m["row_fields"] + [c["label"] for c in m["columns"]] + ["Gap"]
    out = [head]
    for r in m["rows"]:
        out.append([r["label"]] + [r["fields"].get(f, "") for f in m["row_fields"]]
                   + [r["cells"].get(c["key"], "") for c in m["columns"]] + [r.get("gap_note", "")])
    return out


def to_csv(m: dict) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    for line in _table(m):
        w.writerow(line)
    return buf.getvalue()


def to_xlsx(matrices: list[dict]) -> bytes:
    """One sheet per matrix, header frozen, AutoFilter on, gaps highlighted."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    head_fill = PatternFill("solid", fgColor="1F3864")
    gap_fill = PatternFill("solid", fgColor="FCE4D6")
    hit_fill = PatternFill("solid", fgColor="E2EFDA")
    for m in matrices:
        ws = wb.create_sheet(m["name"][:31])
        table = _table(m)
        for row in table:
            ws.append(row)
        n_fields = 1 + len(m["row_fields"])
        n_cols = len(table[0])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = head_fill
            cell.alignment = Alignment(wrap_text=True, vertical="top",
                                       text_rotation=90 if cell.column > n_fields and cell.column < n_cols
                                       and m["name"] in ("uc_state", "state_logic", "uc_uc") else 0)
        for i, r in enumerate(m["rows"], start=2):
            for j, col in enumerate(m["columns"], start=n_fields + 1):
                if r["cells"].get(col["key"]):
                    ws.cell(i, j).fill = hit_fill
                elif col["key"] in r.get("gap_cols", []):
                    ws.cell(i, j).fill = gap_fill
            if r["gap"]:
                ws.cell(i, n_cols).fill = gap_fill
        ws.freeze_panes = ws.cell(2, n_fields + 1)
        ws.auto_filter.ref = f"A1:{get_column_letter(n_cols)}{max(1, len(table))}"
        for j in range(1, n_cols + 1):
            width = max(len(str(ws.cell(i, j).value or "")) for i in range(2, len(table) + 1)) if len(table) > 1 else 8
            ws.column_dimensions[get_column_letter(j)].width = min(max(6, width + 2), 45 if j <= n_fields else 30)
        ws.append([])
        ws.append([m["legend"]])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
